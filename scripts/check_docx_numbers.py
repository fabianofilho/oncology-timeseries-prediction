#!/usr/bin/env python3
"""Confere o pacote da CSP contra paper/verified_numbers.json, numero por numero.

Passo obrigatorio antes de entregar: nao e feito de cabeca, e feito por este
comparador. Le de volta os quatro .docx do pacote (paragrafos e tabelas), procura
cada valor que o JSON diz e reporta o que falta.

Alguns numeros ficaram fora do pacote de proposito, no corte de 10 para 5
ilustracoes (CORTADAS em build_csp_submission.py). Eles aparecem aqui com o motivo,
para o corte ser visivel e nao ser confundido com esquecimento.

Tambem valida a estrutura que a CSP exige e que o python-docx nao garante sozinho:
numeracao de linha e numero de pagina no texto, tabelas nativas, figuras em arquivo
separado com pelo menos 300 dpi.

Ate 2026-09-28 conferia build/manuscript.docx, o manuscrito de revisao gerado por
scripts/build_docx.py, aposentado nessa data: o pacote da CSP passou a ser a unica
versao do texto.

Uso:
    python scripts/check_docx_numbers.py [--dir build/csp]
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
PAPER = ROOT / "paper"
ARQUIVOS = ("CSP_texto", "CSP_tabelas", "CSP_legendas_figuras", "CSP_material_suplementar")
DPI_MINIMO = 300
N_FIGURAS = 3
N_TABELAS = 2


def texto_completo(doc: Document) -> str:
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            partes.extend(c.text for c in row.cells)
    return "\n".join(partes)


def main() -> None:
    ap = argparse.ArgumentParser(description="Confere numeros do pacote da CSP")
    ap.add_argument("--dir", default="build/csp")
    args = ap.parse_args()
    pasta = ROOT / args.dir

    v = json.loads((PAPER / "verified_numbers.json").read_text(encoding="utf-8"))
    docs = {nome: Document(str(pasta / f"{nome}.docx"))
            for nome in ARQUIVOS if (pasta / f"{nome}.docx").exists()}
    txt = "\n".join(texto_completo(d) for d in docs.values())

    co, ba = v["coorte"], v["benchmark_agregado"]
    ci, met, crit = ba["ci_smape"], ba["metricas"], ba["criterio"]
    ext, tend, sit, reg = v["extrapolacao"], v["tendencia_pre_pandemia"], v["sitios"], v["teste_regime"]

    checagens: list[tuple[str, str, str | None]] = []

    def esperar(rotulo: str, valor: str, fora: str | None = None) -> None:
        """`fora` e o motivo quando o numero ficou fora do pacote por decisao."""
        checagens.append((rotulo, valor, fora))

    esperar("obitos por cancer", f"{co['obitos_por_cancer']:,}")
    esperar("pontos da serie completa", str(co["pontos_serie_completa"]))
    esperar("pontos da serie de treino", str(co["pontos_serie_treino"]))
    esperar("media mensal", f"{int(co['media_mensal']):,}")
    esperar("minimo mensal", f"{co['minimo_mensal']:,}")
    esperar("maximo mensal", f"{co['maximo_mensal']:,}")
    esperar("janelas de rolling origin", str(ba["n_janelas"]))
    esperar("previsoes por modelo", str(ba["n_previsoes_por_modelo"]))
    esperar("min train", str(ba["min_train"]))
    esperar("horizonte", str(ba["horizonte"]))
    esperar("replicas do bootstrap", f"{int(ci['sarima']['B']):,}")
    esperar("pares confirmados", str(crit["pares_confirmados"]))
    esperar("pares totais", str(crit["pares_total"]))
    esperar("alpha do criterio", str(crit["alpha"]))
    esperar("min horizontes do criterio", str(crit["min_horizontes"]))

    for m in ci:
        esperar(f"sMAPE pontual de {m}", f"{ci[m]['point']:.2f}")
        esperar(f"IC de {m}", f"{ci[m]['ci_low']:.2f} to {ci[m]['ci_high']:.2f}")
        esperar(f"MAE de {m}", f"{met[m]['mae']:.1f}")

    esperar("R2 da tendencia pre-pandemia", str(tend["r2"]))
    esperar("slope da tendencia", f"{int(tend['slope_obitos_por_ano']):,}")
    esperar("pct de observacoes acima do max de treino",
            f"{ext['_observacoes_acima_do_max_treino']['pct']}")
    esperar("n de observacoes acima do max de treino",
            str(ext["_observacoes_acima_do_max_treino"]["n"]))
    for m in ("catboost", "xgboost", "sarima", "prophet", "timesfm"):
        esperar(f"extrapolacoes de {m}", str(ext[m]["acima_do_max_treino"]))

    # A tabela por sitio foi cortada: o texto carrega o agregado e a prostata, os
    # extremos, e a Figura 3 carrega os demais.
    no_texto = ("agregado", "prostata")
    for s in ("agregado", "pulmao", "colorretal", "mama", "prostata"):
        fora = None if s in no_texto else "tabela por sitio cortada; valor so na Figura 3"
        esperar(f"sMAPE do sitio {s}", f"{sit[s]['smape_melhor']:.2f}", fora)
        esperar(f"media mensal do sitio {s}", f"{int(sit[s]['media_mensal']):,}", fora)
    esperar("cobertura dos sitios", str(v["sitios_cobertura_pct"]))

    # O teste de regime virou texto com o ingenuo sazonal, o SARIMA e o Prophet inline.
    for m in reg["por_modelo"]:
        fora = (None if m in ("seasonal_naive", "sarima", "prophet")
                else "tabela do teste de regime cortada; so os modelos citados no texto")
        esperar(f"sMAPE fora de regime de {m}", f"{reg['por_modelo'][m]['smape_total']:.2f}", fora)

    anos = list(co["obitos_por_ano"])
    for a, val in co["obitos_por_ano"].items():
        fora = None if a in (anos[0], anos[-1]) else "obitos por ano cortados; so o primeiro e o ultimo na prosa"
        esperar(f"obitos em {a}", f"{val:,}", fora)

    itv = v.get("intervalos")
    if itv:
        fx = itv["faixa_cobertura"]
        for nivel in ("n50", "n80", "n95"):
            esperar(f"cobertura minima {nivel}", f"{fx[nivel][0]:.3f}")
            esperar(f"cobertura maxima {nivel}", f"{fx[nivel][1]:.3f}")
        esperar("pior WIS", f"{itv['pior_wis']['valor']:.1f}")
        esperar("WIS do ingenuo sazonal", f"{itv['wis_seasonal_naive']:.1f}")
        esperar("cobertura 95 fora do regime", f"{itv['cobertura_95_fora']:.3f}")
        esperar("n avaliado por modelo", str(itv["n_avaliado_por_modelo"]))

    rp_path = ROOT / "results" / "reproducao_benchmark.json"
    if rp_path.exists():
        import math
        rp = json.loads(rp_path.read_text(encoding="utf-8"))
        esperar("sMAPE do XGBoost reproduzido", f"{rp['modelos']['xgboost']['smape_reproduzido']:.2f}")
        esperar("tolerancia dos outros seis modelos",
                f"{math.ceil(rp['max_abs_dif_smape_sem_xgboost'] * 1000) / 1000:.3f}")
        for p in rp["pares_que_mudam"]:
            a, b = p["dm_reproduzido"].split("/")
            esperar("DM do par que muda, reproduzido", f"{a} of {b}")

    vd_path = ROOT / "results" / "criterion_verdict_2010_2019_full.json"
    vd = json.loads(vd_path.read_text(encoding="utf-8"))
    if vd["criterio"].get("versao") == 2:
        mudam = sum(p["confirmado"] != p["confirmado_v1"] for p in vd["pares"])
        if mudam == 0:
            esperar("criterio v2 nao muda veredito do agregado", f"none of the {len(vd['pares'])} verdicts")

    si_path = ROOT / "results" / "serie_inteira_resumo.json"
    if si_path.exists():
        si = json.loads(si_path.read_text(encoding="utf-8"))
        esperar("janelas da serie inteira", f"untreated ({si['janelas']} windows)")

    exigidos = [(r, x) for r, x, fora in checagens if fora is None]
    cortados = [(r, x, fora) for r, x, fora in checagens if fora is not None]
    faltando = [(r, x) for r, x in exigidos if x not in txt]
    print(f"Numeros exigidos no pacote: {len(exigidos)}")
    print(f"Encontrados: {len(exigidos) - len(faltando)}")
    if faltando:
        print(f"\nNAO ENCONTRADOS ({len(faltando)}):")
        for r, x in faltando:
            print(f"  ERRO  {r:44s} esperado no texto: {x}")
    else:
        print("Todos os valores exigidos aparecem no pacote com o mesmo valor.")
    print(f"\nFora do pacote por decisao ({len(cortados)}):")
    for motivo in sorted({f for _, _, f in cortados}):
        n = sum(1 for _, _, f in cortados if f == motivo)
        print(f"  {n:>2}  {motivo}")

    print("\n--- estrutura ---")
    faltam_arquivos = [a for a in ARQUIVOS if a not in docs]
    texto = docs.get("CSP_texto")
    sec = texto.sections[0] if texto else None
    tem_ln = bool(sec) and sec._sectPr.find(qn("w:lnNumType")) is not None
    tem_pg = bool(sec) and "PAGE" in sec.footer.paragraphs[0]._p.xml
    tabs = len(docs["CSP_tabelas"].tables) if "CSP_tabelas" in docs else 0
    embutidas = len(texto.inline_shapes) if texto else 0
    figs = sorted((pasta / "figures").glob("*.png"))
    dpis = {f.name: Image.open(f).info.get("dpi", (0, 0))[0] for f in figs}
    # o PNG grava pixels por metro, entao 300 dpi volta como 299,9994
    dpi_ok = all(d >= DPI_MINIMO - 0.5 for d in dpis.values())
    pend = len(re.findall(r"\[PENDING:", txt))

    print(f"arquivos do pacote:                         {'OK' if not faltam_arquivos else 'FALTAM ' + str(faltam_arquivos)}")
    print(f"numeracao de linha no texto (w:lnNumType):  {'OK' if tem_ln else 'AUSENTE'}")
    print(f"campo de numero de pagina no rodape:        {'OK' if tem_pg else 'AUSENTE'}")
    print(f"figuras embutidas no texto (devem ser 0):   {embutidas}")
    print(f"tabelas nativas em CSP_tabelas:             {tabs} (esperado {N_TABELAS})")
    print(f"figuras em arquivo separado:                {len(figs)} (esperado {N_FIGURAS})")
    print(f"dpi das figuras (minimo {DPI_MINIMO}):              "
          + ", ".join(f"{k} {d:.0f}" for k, d in dpis.items()))
    print(f"marcadores [PENDING]:                       {pend}")
    ficha = pasta / "CSP_formulario_submissao.md"
    pend_ficha = (len(re.findall(r"\[PENDING:", ficha.read_text(encoding="utf-8")))
                  if ficha.exists() else None)
    print(f"ficha do formulario (PENDING):              "
          + ("AUSENTE" if pend_ficha is None else str(pend_ficha)))

    ok = (not faltando and not faltam_arquivos and tem_ln and tem_pg and embutidas == 0
          and tabs == N_TABELAS and len(figs) == N_FIGURAS and dpi_ok)
    print("\nVEREDITO: " + ("estrutura e numeros OK" if ok else "REVISAR os itens acima")
          + (f" ({pend} PENDING a resolver antes de submeter)" if pend else "")
          + (f" (ficha do formulario: {pend_ficha} PENDING)" if pend_ficha else ""))
    raise SystemExit(0 if ok else 1)


if __name__ == "__main__":
    main()
