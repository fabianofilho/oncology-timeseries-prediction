#!/usr/bin/env python3
"""Comunicacao Breve para a CSP: mortalidade por cancer de pulmao segundo sexo (LAB-182).

Gera build/csp_sexo/: CB_texto.docx, CB_tabela.docx, CB_legenda.docx e figures/Figure1.png
(300 dpi). Formato da Comunicacao Breve da CSP: ate 2.500 palavras e 3 ilustracoes, com
Resumo, Abstract e Resumen de ate 250 palavras cada. Aqui sao 2 ilustracoes (1 tabela e
1 figura).

Como no manuscrito principal, nenhum numero e digitado: tudo sai de results/
(sexo_tendencia_pulmao.json, sexo_taxas_pulmao.csv, sex_series_report.json e os
benchmarks por sexo). Citacao por chave [@chave], numerada pela ordem de aparicao.

Uso:
    python scripts/build_sexo_cb.py [--out build/csp_sexo]
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import re
from pathlib import Path

import matplotlib
import matplotlib.ticker

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
R = ROOT / "results"

# formatacao compartilhada com o pacote principal
_spec = importlib.util.spec_from_file_location("csp", ROOT / "scripts" / "build_csp_submission.py")
csp = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(csp)

LIMITE_PALAVRAS = 2500
LIMITE_ILUSTRACOES = 3
LIMITE_RESUMO = 250

REFERENCES = {
    "reis2026": "Reis RS, Lima FCDS, Silva DHND, Scaff AJM, Szklo AS. Projections for lung "
    "cancer mortality in Brazil by 2045: a regional analysis. Cad Saúde Pública 2026; "
    "42:e00199725.",
    "malta2016": "Malta DC, Abreu DM, Moura L, Lana GC, Azevedo G, França E. Trends in "
    "corrected lung cancer mortality rates in Brazil and regions. Rev Saúde Pública 2016; "
    "50:33.",
    "souza2012": "Souza MC, Vasconcelos AG, Cruz OG. Trends in lung cancer mortality in Brazil "
    "from the 1980s into the early 21st century: age-period-cohort analysis. Cad Saúde "
    "Pública 2012; 28:21-30.",
    "malta2017": "Malta DC, Stopa SR, Santos MAS, Andrade SSCA, Oliveira TP, Cristo EB, et al. "
    "Evolution of tobacco use indicators according to telephone surveys, 2006-2014. Cad Saúde "
    "Pública 2017; 33 Suppl 3:e00134915.",
    "lima2014": "Lima EEC, Queiroz BL. Evolution of the deaths registry system in Brazil: "
    "associations with changes in the mortality profile, under-registration of death counts, "
    "and ill-defined causes of death. Cad Saúde Pública 2014; 30:1721-30.",
    "doll1966": "Doll R, Payne P, Waterhouse JAH, editors. Cancer incidence in five "
    "continents. Vol. I. Geneva: Union Internationale Contre le Cancer; 1966.",
    "fernandes2020": "Fernandes GA, Menezes FDS, Silva LF, Antunes JLF, Toporcov TN. "
    "Inequalities in lung cancer mortality trends in Brazil, 2000-2015. Sci Rep 2020; "
    "10:19164.",
}
CITE = re.compile(r" ?\[@([a-z0-9]+(?:,@[a-z0-9]+)*)\]")
CITADAS: list[str] = []

ORDEM = ["seasonal_naive", "linear_trend", "sarima", "prophet", "timesfm", "xgboost", "catboost"]
NOME_PT = {"seasonal_naive": "ingênuo sazonal", "linear_trend": "tendência linear",
           "sarima": "SARIMA", "prophet": "Prophet", "timesfm": "TimesFM",
           "xgboost": "XGBoost", "catboost": "CatBoost"}


# ------------------------------------------------------------------ numeros
def br(x: float, dec: int = 2) -> str:
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def en(x: float, dec: int = 2) -> str:
    return f"{x:,.{dec}f}"


def lista_pt(itens: list[str]) -> str:
    return itens[0] if len(itens) == 1 else ", ".join(itens[:-1]) + " e " + itens[-1]


def carrega() -> dict:
    tend = json.loads((R / "sexo_tendencia_pulmao.json").read_text(encoding="utf-8"))
    if tend["veredito"]["hipotese"] != "apoiada":
        raise SystemExit("[ERRO] o texto foi escrito para a hipotese apoiada; resultado mudou")
    T = {(t["sexo"], t["janela"], t["medida"]): t for t in tend["tendencias"]}
    tx = pd.read_csv(R / "sexo_taxas_pulmao.csv")
    ser = json.loads((R / "sex_series_report.json").read_text(encoding="utf-8"))["desfechos"]["pulmao"]
    pop = pd.read_csv(R / "populacao_sp_sexo_faixa_2010_2023.csv")
    bench = {}
    for sx in ("masc", "fem"):
        b = pd.read_csv(R / f"uncertainty_onco_SP_sexo_pulmao_{sx}_2010_2019_bootstrap_metrics.csv")
        sm = b[b.metric == "smape"].set_index("model").point
        met = pd.read_csv(R / f"benchmark_onco_SP_sexo_pulmao_{sx}_2010_2019_metrics.csv").set_index("model")
        vd = json.loads((R / f"criterion_verdict_sexo_pulmao_{sx}_2010_2019.json").read_text(encoding="utf-8"))
        if vd["criterio"].get("versao") != 2:
            raise SystemExit("[ERRO] vereditos por sexo fora do criterio v2")
        batem = []
        for p in vd["pares"]:
            a, c = p["par"].split("_menos_")
            outro = c if a == "seasonal_naive" else a if c == "seasonal_naive" else None
            if outro and p["confirmado"] and met.loc[outro, "smape"] < met.loc["seasonal_naive", "smape"]:
                batem.append(outro)
        bench[sx] = {"smape": sm, "melhor": sm.idxmin(), "batem": sorted(batem, key=ORDEM.index),
                     "n_jan": int(b.n_windows.iloc[0]), "media": ser["series"][sx]["media_mensal_pre"]}
    taxa = tx.pivot(index="ano", columns="sexo", values="taxa_padronizada")
    obitos = tx.pivot(index="ano", columns="sexo", values="obitos")
    return {"T": T, "tx": tx, "taxa": taxa, "obitos": obitos, "ser": ser, "pop": pop,
            "bench": bench, "razao": (taxa.fem / taxa.masc), "ign_pct": tend["idade_ignorada_pct"]}


def apc(t: dict, lingua: str) -> str:
    f = br if lingua in ("pt", "es") else en
    sep = " a " if lingua in ("pt", "es") else " to "
    rot = "IC95%" if lingua in ("pt", "es") else "95%CI"
    return f"{f(t['apc'])}% ({rot}: {f(t['ic95'][0])}{sep}{f(t['ic95'][1])})"


# ------------------------------------------------------------------ texto
def cite_runs(p, text: str) -> None:
    pos = 0
    for m in CITE.finditer(text):
        p.add_run(text[pos:m.start()])
        nums = []
        for k in m.group(1).split(",@"):
            if k not in REFERENCES:
                raise KeyError(f"citacao sem referencia: {k}")
            if k not in CITADAS:
                CITADAS.append(k)
            nums.append(str(CITADAS.index(k) + 1))
        p.add_run(",".join(nums)).font.superscript = True
        pos = m.end()
    p.add_run(text[pos:])


def par(doc, text: str, lead: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if lead:
        p.add_run(lead).bold = True
    cite_runs(p, text)
    return p


def resumos(D: dict) -> list[tuple[str, str, str, str]]:
    T, ser, raz = D["T"], D["ser"], D["razao"]
    f19m, f19f = T[("masc", "2010-2019", "taxa_padronizada")], T[("fem", "2010-2019", "taxa_padronizada")]
    c19m, c19f = T[("masc", "2010-2019", "obitos")], T[("fem", "2010-2019", "obitos")]
    s23m, s23f = T[("masc", "2010-2023", "taxa_padronizada")], T[("fem", "2010-2023", "taxa_padronizada")]
    a0, a1 = int(raz.index.min()), int(raz.index.max())
    n = ser["obitos"]
    pt = (
        f"A mortalidade por câncer de pulmão no Brasil vem caindo entre homens e subindo entre "
        f"mulheres, mas o número de óbitos também responde ao envelhecimento da população. "
        f"Analisamos os {br(n, 0)} óbitos por câncer de traqueia, brônquios e pulmão de "
        f"residentes no Estado de São Paulo entre {a0} e {a1}, registrados no Sistema de "
        f"Informações sobre Mortalidade. Calculamos taxas anuais padronizadas por idade pela "
        f"população mundial e estimamos a variação percentual anual por regressão de "
        f"Prais-Winsten, com 2010 a 2019 como período principal e a regra de decisão declarada "
        f"antes da análise. Nesse período, a taxa padronizada caiu entre homens, "
        f"{apc(f19m, 'pt')} ao ano, e subiu entre mulheres, {apc(f19f, 'pt')}, enquanto o "
        f"número de óbitos cresceu nos dois sexos, {br(c19m['apc'])}% e {br(c19f['apc'])}% ao "
        f"ano. A razão entre as taxas feminina e masculina passou de {br(raz.loc[a0])} em {a0} "
        f"para {br(raz.loc[a1])} em {a1}. Incluindo os anos da pandemia, a queda masculina se "
        f"manteve, {apc(s23m, 'pt')}, e a alta feminina deixou de ser significativa, "
        f"{apc(s23f, 'pt')}. A taxa mede o risco e a contagem mede a demanda: com a população "
        f"envelhecendo, os óbitos crescem mesmo onde o risco cai, e a distância entre os sexos "
        f"se estreita."
    )
    en_ = (
        f"Lung cancer mortality in Brazil has been falling among men and rising among women, but "
        f"the number of deaths also responds to population ageing. We analysed the "
        f"{en(n, 0)} deaths from cancer of the trachea, bronchus and lung among residents of São "
        f"Paulo State from {a0} to {a1}, recorded in the Mortality Information System. We "
        f"computed annual rates age-standardised to the world population and estimated the "
        f"annual percent change by Prais-Winsten regression, with 2010 to 2019 as the main "
        f"period and the decision rule declared before the analysis. Over that period the "
        f"age-standardised rate fell among men, {apc(f19m, 'en')} per year, and rose among women, "
        f"{apc(f19f, 'en')}, while the number of deaths grew in both sexes, {en(c19m['apc'])}% "
        f"and {en(c19f['apc'])}% per year. The female-to-male rate ratio went from "
        f"{en(raz.loc[a0])} in {a0} to {en(raz.loc[a1])} in {a1}. Including the pandemic years, "
        f"the decline among men held, {apc(s23m, 'en')}, and the rise among women was no longer "
        f"significant, {apc(s23f, 'en')}. Rates measure risk and counts measure demand: in an "
        f"ageing population, deaths grow even where risk falls, and the gap between the sexes is "
        f"narrowing."
    )
    es = (
        f"La mortalidad por cáncer de pulmón en Brasil viene disminuyendo entre hombres y "
        f"aumentando entre mujeres, pero el número de muertes también responde al envejecimiento "
        f"de la población. Analizamos las {br(n, 0)} muertes por cáncer de tráquea, bronquios y "
        f"pulmón de residentes del Estado de São Paulo entre {a0} y {a1}, registradas en el "
        f"Sistema de Información sobre Mortalidad. Calculamos tasas anuales estandarizadas por "
        f"edad con la población mundial y estimamos la variación porcentual anual mediante "
        f"regresión de Prais-Winsten, con 2010 a 2019 como período principal y la regla de "
        f"decisión declarada antes del análisis. En ese período, la tasa estandarizada disminuyó "
        f"entre hombres, {apc(f19m, 'es')} al año, y aumentó entre mujeres, {apc(f19f, 'es')}, "
        f"mientras el número de muertes creció en ambos sexos, {br(c19m['apc'])}% y "
        f"{br(c19f['apc'])}% al año. La razón entre las tasas femenina y masculina pasó de "
        f"{br(raz.loc[a0])} en {a0} a {br(raz.loc[a1])} en {a1}. Al incluir los años de la "
        f"pandemia, la caída masculina se mantuvo, {apc(s23m, 'es')}, y el aumento femenino dejó "
        f"de ser significativo, {apc(s23f, 'es')}. La tasa mide el riesgo y el conteo mide la "
        f"demanda: con una población que envejece, las muertes crecen aun donde el riesgo "
        f"disminuye, y la distancia entre los sexos se estrecha."
    )
    return [
        ("Resumo", pt, "Neoplasias Pulmonares; Mortalidade; Fatores Sexuais; Sistemas de "
         "Informação", "Palavras-chave: "),
        ("Abstract", en_, "Lung Neoplasms; Mortality; Sex Factors; Information Systems",
         "Keywords: "),
        ("Resumen", es, "Neoplasias Pulmonares; Mortalidad; Factores Sexuales; Sistemas de "
         "Información", "Palabras clave: "),
    ]


def build_texto(D: dict):
    CITADAS.clear()
    T, ser, raz, taxa, ob, pop, B = (D["T"], D["ser"], D["razao"], D["taxa"], D["obitos"],
                                     D["pop"], D["bench"])
    a0, a1 = int(taxa.index.min()), int(taxa.index.max())
    doc = csp.base_doc()
    p = doc.add_paragraph()
    p.add_run(
        "Mortalidade por câncer de pulmão segundo sexo no Estado de São Paulo, Brasil, "
        f"{a0}-{a1}: taxas padronizadas em sentidos opostos e óbitos em alta nos dois sexos"
    ).bold = True
    csp.flag_par(doc, "autores, afiliações, ORCID, autor correspondente e contribuições")

    for titulo, corpo, chaves, rot in resumos(D):
        csp.h(doc, titulo)
        par(doc, corpo)
        q = doc.add_paragraph()
        q.add_run(rot).bold = True
        q.add_run(chaves)

    # ---- Introducao
    csp.h(doc, "Introdução")
    par(doc,
        "Entre 1996 e 2011, a mortalidade por câncer de pulmão corrigida para sub-registro e "
        "causas mal definidas caiu entre homens e subiu entre mulheres no Brasil [@malta2016]. "
        "Análises de idade, período e coorte associam esse padrão à história do tabagismo e às "
        "medidas de controle do tabaco adotadas desde 1986, com risco decrescente nas coortes "
        "masculinas mais jovens e crescente entre as femininas [@souza2012]. Projeções publicadas "
        "nesta revista antecipam mais óbitos até 2045 mesmo com taxas ajustadas em queda, porque "
        "o envelhecimento eleva as taxas brutas [@reis2026].")
    par(doc,
        "Taxa padronizada e número de óbitos respondem a perguntas diferentes. A primeira "
        "estima o risco livre da composição etária; o segundo é a demanda que chega aos "
        "serviços. Numa população que envelhece, as duas podem andar em sentidos opostos, e "
        "uma conclusão tirada de uma delas não vale para a outra.")
    par(doc,
        f"Descrevemos a tendência da mortalidade por câncer de pulmão segundo sexo no Estado de "
        f"São Paulo de {a0} a {a1}, contrastando taxa padronizada e contagem, com o critério "
        f"de apoio à hipótese declarado antes da análise.")

    # ---- Metodos
    csp.h(doc, "Métodos")
    par(doc,
        f"Usamos os óbitos do Sistema de Informações sobre Mortalidade por residência no Estado "
        f"de São Paulo, de {a0} a {a1}, com causa básica de câncer de traqueia, brônquios e "
        f"pulmão (CID-10 C33-C34). A completude do registro de óbitos no Sudeste está entre as "
        f"mais altas do país [@lima2014]. Registros sem sexo informado ({br(ser['sexo_ignorado'], 0)}) "
        f"foram excluídos. A idade foi decodificada do campo de idade do sistema; óbitos sem "
        f"idade ficaram fora do numerador e seriam redistribuídos proporcionalmente se "
        f"passassem de 1% do total.")
    par(doc,
        f"Os denominadores vieram das estimativas populacionais por município, idade e sexo "
        f"2000-2025 do Ministério da Saúde e da Rede Interagencial de Informações para a Saúde, "
        f"em faixas de cinco anos até 80 anos ou mais. As taxas anuais foram padronizadas pelo "
        f"método direto, com a população mundial de Segi modificada por Doll [@doll1966]; como "
        f"o denominador fecha em 80 anos ou mais, os pesos das faixas de 80 a 84 e de 85 ou mais "
        f"foram somados.")
    par(doc,
        "A tendência foi estimada por regressão de Prais-Winsten do logaritmo decimal da taxa "
        "contra o ano, que acomoda autocorrelação de primeira ordem nos resíduos [@fernandes2020]; "
        "a variação percentual anual (VPA) é (10^b - 1) x 100, com intervalo de confiança de 95% "
        "pela distribuição t. O período principal foi 2010 a 2019, antes da pandemia de COVID-19; "
        f"2010 a {a1} foi a análise de sensibilidade. A mesma regressão sobre o número de óbitos "
        "serviu de contraste, para mostrar quanto do movimento vem da composição etária.")
    par(doc,
        "A regra de decisão foi registrada no rastreador de tarefas do projeto antes de qualquer "
        "análise: a hipótese de que a mortalidade cai entre homens e sobe entre mulheres só "
        "seria considerada apoiada se, no período principal, a VPA masculina fosse negativa e a "
        "feminina positiva, ambas com intervalo de confiança excluindo zero. Uma cópia datada "
        "desse registro está no repositório do projeto.")
    par(doc,
        f"Como análise complementar, as séries mensais de óbitos de cada sexo foram submetidas "
        f"a sete métodos de previsão (do ingênuo sazonal, que repete o mesmo mês do ano "
        f"anterior, a um modelo de fundação), com validação por origem móvel de 2010 a 2019, "
        f"horizonte de seis meses e {B['masc']['n_jan']} janelas. Uma diferença entre métodos "
        f"só foi considerada confirmada se o intervalo de confiança da diferença, por "
        f"reamostragem de janelas inteiras, excluísse zero tanto no erro percentual absoluto "
        f"simétrico (sMAPE) quanto no erro absoluto médio, e o teste de Diebold-Mariano "
        f"apontasse diferença em ao menos três dos seis horizontes.")
    par(doc,
        "O código e as séries agregadas estão em "
        "https://github.com/fabianofilho/oncology-timeseries-prediction. Os dados são públicos e "
        "sem identificação e, conforme a Resolução nº 510/2016 do Conselho Nacional de Saúde, a "
        "pesquisa dispensa apreciação por comitê de ética.")

    # ---- Resultados
    f19m, f19f = T[("masc", "2010-2019", "taxa_padronizada")], T[("fem", "2010-2019", "taxa_padronizada")]
    c19m, c19f = T[("masc", "2010-2019", "obitos")], T[("fem", "2010-2019", "obitos")]
    s23m, s23f = T[("masc", f"2010-{a1}", "taxa_padronizada")], T[("fem", f"2010-{a1}", "taxa_padronizada")]
    c23m, c23f = T[("masc", f"2010-{a1}", "obitos")], T[("fem", f"2010-{a1}", "obitos")]
    # o texto foi escrito para este padrao de resultado; se ele mudar, o build para
    raz_ano = raz.sort_index()
    if not (raz_ano.iloc[-1] > raz_ano.iloc[0] and not raz_ano.is_monotonic_increasing):
        raise SystemExit("[ERRO] frase sobre o estreitamento da razao F/M precisa ser revista")
    if not (s23m["apc"] < f19m["apc"] < 0 and s23m["sentido"] == "queda"
            and s23f["sentido"] == "estavel" and c19m["sentido"] == c19f["sentido"] == "aumento"
            and c23m["sentido"] == c23f["sentido"] == "aumento"):
        raise SystemExit("[ERRO] frases da sensibilidade e do contraste precisam ser revistas")
    pop0 = pop[pop.ano == a0].groupby("sexo").populacao.sum()
    pop1 = pop[pop.ano == a1].groupby("sexo").populacao.sum()
    csp.h(doc, "Resultados")
    par(doc,
        f"Foram {br(ser['obitos'], 0)} óbitos por câncer de pulmão no período, "
        f"{br(ser['series']['masc']['obitos'], 0)} entre homens e "
        f"{br(ser['series']['fem']['obitos'], 0)} entre mulheres; {br(ser['idade_ignorada'], 0)} "
        f"({br(ser['idade_ignorada_pct'], 2)}%) não tinham idade informada, abaixo do limiar que "
        f"exigiria redistribuição. A população do estado passou de "
        f"{br((pop0.sum()) / 1e6, 1)} para {br((pop1.sum()) / 1e6, 1)} milhões de habitantes "
        f"entre {a0} e {a1}.")
    par(doc,
        f"A taxa padronizada masculina foi de {br(taxa.loc[a0, 'masc'])} por 100 mil em {a0} "
        f"para {br(taxa.loc[2019, 'masc'])} em 2019 e {br(taxa.loc[a1, 'masc'])} em {a1}; a "
        f"feminina, de {br(taxa.loc[a0, 'fem'])} para {br(taxa.loc[2019, 'fem'])} e "
        f"{br(taxa.loc[a1, 'fem'])} (Figura 1). No período principal, a VPA foi de "
        f"{apc(f19m, 'pt')} entre homens e {apc(f19f, 'pt')} entre mulheres (Tabela 1), o que "
        f"atende ao critério declarado: a hipótese é apoiada. A razão entre as taxas feminina e "
        f"masculina passou de {br(raz.loc[a0])} em {a0} para {br(raz.loc[2019])} em 2019 e "
        f"{br(raz.loc[a1])} em {a1}.")
    par(doc,
        f"O número de óbitos cresceu nos dois sexos, inclusive entre homens, cuja taxa caía: "
        f"VPA de {apc(c19m, 'pt')} para homens e {apc(c19f, 'pt')} para mulheres. Entre "
        f"{a0} e 2019, os óbitos masculinos passaram de {br(ob.loc[a0, 'masc'], 0)} para "
        f"{br(ob.loc[2019, 'masc'], 0)} por ano, e os femininos, de {br(ob.loc[a0, 'fem'], 0)} "
        f"para {br(ob.loc[2019, 'fem'], 0)}.")
    par(doc,
        f"Incluindo os anos da pandemia, a queda masculina se acentuou, com VPA de "
        f"{apc(s23m, 'pt')}, e a alta feminina deixou de ser significativa, {apc(s23f, 'pt')}. "
        f"Os óbitos seguiram em alta nos dois sexos ({br(c23m['apc'])}% e {br(c23f['apc'])}% ao "
        f"ano).")
    mm, mf = B["masc"], B["fem"]
    par(doc,
        f"Na previsão mensal, o erro do melhor método foi de {br(mm['smape'][mm['melhor']])}% "
        f"entre homens ({NOME_PT[mm['melhor']]}, média de {br(mm['media'], 0)} óbitos por mês) "
        f"e de {br(mf['smape'][mf['melhor']])}% entre mulheres ({NOME_PT[mf['melhor']]}, "
        f"{br(mf['media'], 0)} por mês). Bateram o ingênuo sazonal de forma confirmada "
        f"{lista_pt([NOME_PT[x] for x in mm['batem']])} entre homens, e "
        f"{lista_pt([NOME_PT[x] for x in mf['batem']])} entre mulheres; XGBoost e CatBoost não "
        f"o fizeram em nenhum dos sexos.")

    # ---- Discussao
    csp.h(doc, "Discussão")
    par(doc,
        "No Estado de São Paulo, de 2010 a 2019, a mortalidade padronizada por câncer de "
        "pulmão caiu entre homens e subiu entre mulheres, o mesmo padrão descrito para o Brasil "
        "entre 1996 e 2011 [@malta2016], agora com o critério de apoio fixado antes da análise. "
        "A distância entre os sexos se estreitou entre o início e o fim do período, ainda que "
        "não a cada ano. O resultado é "
        "compatível com a história do tabagismo por coorte [@souza2012] e com a queda da "
        "prevalência de fumantes nos dois sexos desde 2006 [@malta2017], mas os dados de "
        "mortalidade não permitem atribuir a tendência a essa causa.")
    vezes = c19f["apc"] / f19f["apc"]
    par(doc,
        "O achado de maior consequência prática é a divergência entre taxa e contagem. Entre "
        "homens, o risco caiu e o número de óbitos subiu; entre mulheres, os dois subiram, mas "
        f"a VPA da contagem foi {br(vezes, 1)} vezes a da taxa. O planejamento de leitos, "
        "oncologia torácica e cuidados paliativos responde à contagem, e uma leitura apenas "
        "das taxas subestimaria a demanda. O mesmo vale para as projeções nacionais, que "
        "antecipam mais óbitos com taxas ajustadas em queda [@reis2026].")
    par(doc,
        "Este trabalho tem limitações. A análise cobre um estado, com registro de óbitos "
        "melhor que o da média do país [@lima2014], e não corrige causas mal definidas, como "
        "fez a série nacional [@malta2016]. As estimativas populacionais de 2024 substituem "
        "séries anteriores e diferem delas. A faixa etária aberta em 80 anos reúne dois pesos "
        "da população padrão, o que pesa mais no câncer de pulmão, concentrado em idade "
        "avançada. Os anos de 2020 a 2023 misturam tendência e efeito da pandemia; por isso o "
        "período principal termina em 2019. Trata-se de um estudo ecológico sem dados de "
        "exposição.")
    par(doc,
        "A queda do risco masculino e a alta do feminino seguem em São Paulo até o fim da "
        "década passada, e a demanda segue crescendo nos dois sexos. Medir só uma das duas "
        "coisas leva a conclusões opostas sobre o mesmo câncer.")

    # ---- Informacoes adicionais
    csp.h(doc, "Informações adicionais")
    par(doc,
        "Os autores usaram um assistente de programação com inteligência artificial (Claude, "
        "Anthropic) para implementar o código de análise e redigir o texto. A ferramenta não é "
        "autora. Os autores revisaram todo o resultado, conferiram cada número contra o código "
        "e assumem a responsabilidade pelo conteúdo.", lead="Uso de inteligência artificial. ")
    par(doc, "Não houve participação de pacientes ou do público neste estudo.",
        lead="Participação de pacientes e do público. ")
    for titulo in ("Colaboradores. ", "Financiamento. ", "Conflito de interesses. ",
                   "Agradecimentos. "):
        q = doc.add_paragraph()
        q.add_run(titulo).bold = True
        csp.flag(q, titulo.strip().rstrip(".").lower())

    csp.h(doc, "Referências")
    sobra = set(REFERENCES) - set(CITADAS)
    if sobra:
        raise ValueError(f"referencias nunca citadas: {sorted(sobra)}")
    for i, k in enumerate(CITADAS, start=1):
        q = doc.add_paragraph()
        q.add_run(f"{i}. {REFERENCES[k]}")

    csp.add_line_numbers(doc.sections[0])
    csp.add_page_numbers(doc.sections[0])
    palavras = sum(len(x.text.split()) for x in doc.paragraphs)
    return doc, palavras


# ------------------------------------------------------------------ tabela, figura, legenda
def build_tabela(D: dict):
    doc = csp.base_doc(tamanho=9, entrelinhas=1.0)
    p = doc.add_paragraph()
    r = p.add_run("Tabela 1")
    r.bold = True
    r.font.size = Pt(10)
    doc.add_paragraph().add_run(
        "Variação percentual anual (VPA) da mortalidade por câncer de pulmão segundo sexo, "
        "taxa padronizada por idade e número de óbitos, Estado de São Paulo, Brasil."
    ).font.size = Pt(9)
    rot_med = {"taxa_padronizada": "Taxa padronizada", "obitos": "Número de óbitos"}
    rot_sent = {"aumento": "aumento", "queda": "queda", "estavel": "estável"}
    linhas = []
    for sx, nome in (("masc", "Masculino"), ("fem", "Feminino")):
        for jan in sorted({k[1] for k in D["T"]}):
            for med in ("taxa_padronizada", "obitos"):
                t = D["T"][(sx, jan, med)]
                linhas.append([nome, jan, rot_med[med], br(t["apc"]),
                               f"{br(t['ic95'][0])} a {br(t['ic95'][1])}", rot_sent[t["sentido"]]])
    csp.tabela(doc, ["Sexo", "Período", "Medida", "VPA (%)", "IC95%", "Tendência"], linhas)
    doc.add_paragraph().add_run(
        "VPA estimada por regressão de Prais-Winsten do logaritmo decimal contra o ano; IC95%: "
        "intervalo de confiança de 95%. Taxa padronizada pelo método direto com a população "
        "mundial de Segi modificada por Doll. Período principal 2010-2019, declarado antes da "
        "análise; 2010-2023 é sensibilidade."
    ).font.size = Pt(8)
    return doc


def build_figura(D: dict, destino: Path) -> None:
    taxa, ob = D["taxa"], D["obitos"]
    fig, (a, b) = plt.subplots(1, 2, figsize=(9, 3.6), dpi=300)
    for ax in (a, b):
        ax.axvspan(2019.5, taxa.index.max() + 0.5, color="0.92", zorder=0)
        ax.spines[["top", "right"]].set_visible(False)
    for sx, rot, cor, mk in (("masc", "Masculino", "#1f4e79", "o"), ("fem", "Feminino", "#b23a48", "s")):
        a.plot(taxa.index, taxa[sx], marker=mk, color=cor, label=rot, lw=1.5, ms=4)
        b.plot(ob.index, ob[sx], marker=mk, color=cor, label=rot, lw=1.5, ms=4)
    a.set_title("A. Taxa padronizada por idade", loc="left", fontsize=10)
    a.set_ylabel("Óbitos por 100 mil habitantes")
    b.set_title("B. Número de óbitos", loc="left", fontsize=10)
    b.set_ylabel("Óbitos por ano")
    a.set_ylim(bottom=0)
    b.set_ylim(bottom=0)
    a.legend(frameon=False, fontsize=8)
    # separador de milhar em portugues no eixo da contagem
    b.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: br(v, 0)))
    for ax in (a, b):
        ax.set_xticks(list(taxa.index[::2]))
        ax.tick_params(labelsize=8)
    fig.tight_layout()
    destino.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(destino, dpi=300)
    plt.close(fig)


def build_legenda(D: dict):
    a0, a1 = int(D["taxa"].index.min()), int(D["taxa"].index.max())
    doc = csp.base_doc()
    p = doc.add_paragraph()
    p.add_run("Legenda da figura").bold = True
    q = doc.add_paragraph()
    q.add_run("Figura 1. ").bold = True
    q.add_run(
        f"Mortalidade por câncer de pulmão (CID-10 C33-C34) segundo sexo, Estado de São "
        f"Paulo, Brasil, {a0}-{a1}. A: taxa anual padronizada por idade pela população mundial "
        f"de Segi modificada por Doll. B: número anual de óbitos. A área sombreada marca 2020 a "
        f"{a1}, fora do período principal da análise de tendência.")
    return doc


def main() -> None:
    ap = argparse.ArgumentParser(description="Comunicacao Breve, cancer de pulmao por sexo")
    ap.add_argument("--out", default="build/csp_sexo")
    args = ap.parse_args()
    out = ROOT / args.out
    D = carrega()
    doc, palavras = build_texto(D)
    out.mkdir(parents=True, exist_ok=True)
    doc.save(str(out / "CB_texto.docx"))
    build_tabela(D).save(str(out / "CB_tabela.docx"))
    build_legenda(D).save(str(out / "CB_legenda.docx"))
    build_figura(D, out / "figures" / "Figure1.png")
    print(f"Pacote em {out}")
    print("Seção a escolher no formulário da CSP: Comunicação Breve")
    print(f"{'palavras (corpo + referencias)':<38}{palavras:>6}   limite {LIMITE_PALAVRAS}"
          f"   {'OK' if palavras <= LIMITE_PALAVRAS else 'EXCEDE'}")
    print(f"{'ilustracoes':<38}{2:>6}   limite {LIMITE_ILUSTRACOES}   OK")
    for titulo, corpo, _, _ in resumos(D):
        n = len(corpo.split())
        print(f"{'palavras do ' + titulo:<38}{n:>6}   limite {LIMITE_RESUMO}"
              f"   {'OK' if n <= LIMITE_RESUMO else 'EXCEDE'}")


if __name__ == "__main__":
    main()
